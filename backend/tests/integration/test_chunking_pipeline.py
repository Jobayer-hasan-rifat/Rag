import hashlib
import io
import logging
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient

from app.core.chunking.types import CHUNKING_VERSION
from tests.conftest import ClientFactory
from tests.files import make_user, markdown_bytes, upload_ok
from tests.helpers import db_execute, db_rows, db_scalar
from tests.pdf_factory import make_pdf
from tests.processing_helpers import (
    chunks_of,
    fetch_document,
    process_now,
    rechunk_now,
    row_of,
    sections_of,
)

pytestmark = pytest.mark.integration

BANGLA = "বাংলাদেশের রাজধানী ঢাকা। এটি একটি পরীক্ষামূলক বাক্য। "


def _upload(client: TestClient, headers: dict[str, str], name: str, data: bytes) -> str:
    return str(upload_ok(client, headers, filename=name, content=data, content_type=None)["id"])


def _long_text(paragraph_count: int = 30) -> bytes:
    return "\n\n".join(
        f"Paragraph {i}. " + "The quick brown fox jumps over the lazy dog. " * 6
        for i in range(paragraph_count)
    ).encode()


def _assert_chunks_are_consistent(db_url: str, doc_id: str) -> list[dict[str, Any]]:
    chunks = chunks_of(db_url, doc_id)
    assert chunks
    assert [c["chunk_index"] for c in chunks] == list(range(len(chunks)))
    for c in chunks:
        text = c["text"]
        assert text == c["section_text"][c["start_char"] : c["end_char"]]
        assert c["char_count"] == len(text)
        assert c["text_sha256"] == hashlib.sha256(text.encode()).hexdigest()
        assert c["chunking_version"] == CHUNKING_VERSION
    return chunks


def _assert_no_chunks_remain(db_url: str, doc_id: str) -> None:
    assert chunks_of(db_url, doc_id) == []
    count, version = db_rows(
        db_url, "SELECT chunk_count, chunking_version FROM documents WHERE id = :d", d=doc_id
    )[0]
    assert count is None and version is None


# --- basic pipeline ------------------------------------------------------------------------


def test_text_document_is_chunked_and_moves_to_chunked_not_ready(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "chunked"
    chunks = _assert_chunks_are_consistent(migrated_database_url, doc_id)
    document = fetch_document(auth_client, alice, doc_id)
    assert document["status"] == "chunked"
    assert document["chunk_count"] == len(chunks) > 3
    assert document["chunking_version"] == CHUNKING_VERSION
    recorded = row_of(migrated_database_url, doc_id)["metadata"]["chunking"]
    assert recorded["version"] == CHUNKING_VERSION
    assert recorded["max_size"] == auth_client.app.state.settings.chunking_max_chars  # type: ignore[attr-defined]


def test_no_document_reaches_ready_in_this_phase(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.md", markdown_bytes())
    process_now(auth_client, doc_id)
    ready = db_scalar(migrated_database_url, "SELECT count(*) FROM documents WHERE status='ready'")
    assert ready == 0


def test_chunk_size_never_exceeds_the_configured_maximum(
    auth_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = auth_client_factory(chunking_max_chars=300, chunking_overlap_chars=40)
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "long.txt", _long_text(40))

    process_now(client, doc_id)

    chunks = _assert_chunks_are_consistent(migrated_database_url, doc_id)
    assert max(c["char_count"] for c in chunks) <= 300


def test_pdf_chunks_carry_page_numbers_and_bangla_text(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    pdf = make_pdf([BANGLA * 6, "Second page in English with enough words to form a chunk."])
    doc_id = _upload(auth_client, alice, "mixed.pdf", pdf)

    process_now(auth_client, doc_id)

    chunks = _assert_chunks_are_consistent(migrated_database_url, doc_id)
    assert {c["page_number"] for c in chunks} == {1, 2}
    assert "ঢাকা" in chunks[0]["text"]
    assert chunks[0]["section_ordinal"] == 0 and chunks[-1]["section_ordinal"] == 1


def test_markdown_headings_become_heading_paths(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    md = (
        b"# Guide\n\nIntroduction text for the guide, long enough to keep.\n\n"
        b"## Setup\n\nInstall the tool and run the first command here.\n\n"
        b"### Details\n\nDetailed notes about the setup process go here.\n\n"
        b"## Usage\n\nUse it like this, with a few more words added.\n"
    )
    doc_id = _upload(auth_client, alice, "guide.md", md)

    process_now(auth_client, doc_id)

    paths = [tuple(c["heading_path"]) for c in chunks_of(migrated_database_url, doc_id)]
    assert ("Guide",) in paths
    assert ("Guide", "Setup") in paths
    assert ("Guide", "Setup", "Details") in paths
    assert ("Guide", "Usage") in paths


def test_docx_headings_are_used_for_chunk_metadata(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    document = docx.Document()
    document.add_heading("Annual Report", level=1)
    document.add_paragraph("Revenue grew steadily across all regions this year.")
    document.add_heading("Outlook", level=2)
    document.add_paragraph("We expect continued growth through the next fiscal year.")
    buffer = io.BytesIO()
    document.save(buffer)
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "report.docx", buffer.getvalue())

    process_now(auth_client, doc_id)

    chunks = chunks_of(migrated_database_url, doc_id)
    assert any(c["heading"] == "Outlook" for c in chunks)
    assert any(tuple(c["heading_path"]) == ("Annual Report", "Outlook") for c in chunks)


def test_document_with_no_text_fails_and_stores_no_chunks(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "empty.pdf", make_pdf(["", ""]))

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed"
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0


# --- determinism and idempotency -----------------------------------------------------------


def test_same_input_produces_identical_chunks_for_different_documents(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    data = _long_text()
    first = _upload(auth_client, alice, "a.txt", data)
    second = _upload(auth_client, alice, "b.txt", data + b"\n")
    process_now(auth_client, first)
    process_now(auth_client, second)

    def shape(doc_id: str) -> list[tuple[Any, ...]]:
        return [
            (c["text"], c["start_char"], c["end_char"], c["overlap_chars"])
            for c in chunks_of(migrated_database_url, doc_id)
        ]

    assert shape(first) == shape(second)


def test_reprocessing_replaces_chunks_without_duplicates(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())
    process_now(auth_client, doc_id)
    before = chunks_of(migrated_database_url, doc_id)
    db_execute(
        migrated_database_url, "UPDATE documents SET status = 'pending' WHERE id = :d", d=doc_id
    )

    assert process_now(auth_client, doc_id).status == "chunked"

    after = chunks_of(migrated_database_url, doc_id)
    assert [c["text"] for c in after] == [c["text"] for c in before]
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == len(after)
    assert fetch_document(auth_client, alice, doc_id)["chunk_count"] == len(after)


def test_processing_a_chunked_document_again_is_a_no_op(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())
    process_now(auth_client, doc_id)
    before = chunks_of(migrated_database_url, doc_id)

    assert process_now(auth_client, doc_id).status == "skipped"

    assert chunks_of(migrated_database_url, doc_id) == before


def test_rechunk_with_new_settings_replaces_chunks_without_re_extracting(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())
    process_now(auth_client, doc_id)
    original = chunks_of(migrated_database_url, doc_id)
    sections_before = sections_of(migrated_database_url, doc_id)
    attempts = row_of(migrated_database_url, doc_id)["processing_attempts"]

    outcome = rechunk_now(auth_client, doc_id, chunking_max_chars=400, chunking_overlap_chars=50)

    assert outcome.status == "chunked"
    rechunked = _assert_chunks_are_consistent(migrated_database_url, doc_id)
    assert len(rechunked) > len(original)
    assert max(c["char_count"] for c in rechunked) <= 400
    assert sections_of(migrated_database_url, doc_id) == sections_before
    row = row_of(migrated_database_url, doc_id)
    assert row["processing_attempts"] == attempts
    assert row["metadata"]["chunking"]["max_size"] == 400
    assert fetch_document(auth_client, alice, doc_id)["chunk_count"] == len(rechunked)


def test_rechunk_ignores_documents_that_are_not_chunked(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())

    assert rechunk_now(auth_client, doc_id).status == "skipped"
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0


# --- failure handling ----------------------------------------------------------------------


def test_chunk_limit_fails_the_document_with_a_safe_reason_and_stores_nothing(
    auth_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = auth_client_factory(
        chunking_max_chars=100, chunking_overlap_chars=0, chunking_max_chunks=3
    )
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "big.txt", _long_text(40))

    outcome = process_now(client, doc_id)

    assert outcome.status == "failed"
    assert row_of(migrated_database_url, doc_id)["failure_reason"] == "too_many_chunks"
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0
    document = fetch_document(client, alice, doc_id)
    assert document["status"] == "failed" and document["chunk_count"] is None
    assert "chunk" in document["error_message"].lower()


def test_unexpected_chunker_error_marks_failed_without_leaking_details(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.core.chunking.chunker import StructureAwareChunker

    def explode(*args: object, **kwargs: object) -> None:
        raise RuntimeError("secret internal detail /srv/path")

    monkeypatch.setattr(StructureAwareChunker, "chunk", explode)
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed"
    row = row_of(migrated_database_url, doc_id)
    assert row["failure_reason"] == "chunking_failed"
    assert "secret" not in str(row["error_message"]) and "/srv" not in str(row["error_message"])
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0


def test_failed_rechunk_marks_failed_and_manual_retry_recovers(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())
    process_now(auth_client, doc_id)

    outcome = rechunk_now(
        auth_client,
        doc_id,
        chunking_max_chars=100,
        chunking_overlap_chars=0,
        chunking_max_chunks=2,
    )

    assert outcome.status == "failed"
    assert row_of(migrated_database_url, doc_id)["failure_reason"] == "too_many_chunks"
    _assert_no_chunks_remain(migrated_database_url, doc_id)
    assert auth_client.get(f"/api/v1/documents/{doc_id}/chunks", headers=alice).json()["data"] == []
    response = auth_client.post(f"/api/v1/documents/{doc_id}/retry", headers=alice)
    assert response.status_code == 202
    assert process_now(auth_client, doc_id).status == "chunked"
    _assert_chunks_are_consistent(migrated_database_url, doc_id)


def test_chunks_are_written_atomically_if_persisting_fails(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.db.repositories.chunk_repository import ChunkRepository

    original = ChunkRepository.replace_all

    async def fail_after_partial_write(
        self: Any, document_id: Any, version: Any, rows: Any
    ) -> None:
        await original(self, document_id, version, rows[:1])
        raise RuntimeError("database went away")

    monkeypatch.setattr(ChunkRepository, "replace_all", fail_after_partial_write)
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())

    process_now(auth_client, doc_id)

    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0


def test_logs_report_chunk_metadata_but_never_chunk_text(
    auth_client: TestClient, caplog: pytest.LogCaptureFixture
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    marker = "ZEBRA-CONFIDENTIAL-MARKER"
    body = f"{marker} appears in this text. ".encode() * 20
    doc_id = _upload(auth_client, alice, "secret.txt", body)

    with caplog.at_level(logging.DEBUG):
        process_now(auth_client, doc_id)

    assert marker not in caplog.text
    assert any(getattr(record, "chunk_count", None) for record in caplog.records)


def test_document_with_only_punctuation_fails_as_empty_not_as_zero_chunks(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "dashes.txt", b"----- ***** =====\n\n..... !!!!!")

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed"
    assert row_of(migrated_database_url, doc_id)["failure_reason"] == "empty_document"
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0


@pytest.mark.parametrize("method", ["complete_extraction", "complete_chunking"])
def test_result_is_discarded_if_the_document_changed_while_processing(
    auth_client: TestClient,
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
    method: str,
) -> None:
    from app.db.repositories.document_repository import DocumentRepository

    async def document_changed(*args: object, **kwargs: object) -> bool:
        return False

    monkeypatch.setattr(DocumentRepository, method, document_changed)
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())

    outcome = process_now(auth_client, doc_id)

    assert (outcome.status, outcome.reason) == ("skipped", "document_changed")
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_chunks") == 0


def test_rechunk_soft_time_limit_marks_the_document_failed(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from celery.exceptions import SoftTimeLimitExceeded

    from app.workers import document_tasks
    from app.workers.celery_app import create_celery_app

    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())
    process_now(auth_client, doc_id)
    db_execute(
        migrated_database_url, "UPDATE documents SET status='chunking' WHERE id = :d", d=doc_id
    )

    async def interrupted(*args: object, **kwargs: object) -> None:
        raise SoftTimeLimitExceeded

    monkeypatch.setattr(document_tasks, "_rechunk", interrupted)
    create_celery_app(auth_client.app.state.settings, configure_logs=False).set_current()  # type: ignore[attr-defined]

    result = document_tasks.rechunk_document.apply(args=[doc_id])

    assert result.get() == {"status": "failed", "reason": "timeout"}
    row = row_of(migrated_database_url, doc_id)
    assert row["status"] == "failed" and row["failure_reason"] == "timeout"
    _assert_no_chunks_remain(migrated_database_url, doc_id)


def test_requeued_document_does_not_keep_serving_old_chunks(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    import asyncio
    import uuid

    from app.db.repositories.document_repository import DocumentRepository
    from app.db.session import create_engine, create_session_factory

    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "notes.txt", _long_text())
    process_now(auth_client, doc_id)
    assert chunks_of(migrated_database_url, doc_id)
    db_execute(
        migrated_database_url, "UPDATE documents SET status='chunking' WHERE id = :d", d=doc_id
    )

    async def release() -> bool:
        engine = create_engine(auth_client.app.state.settings)  # type: ignore[attr-defined]
        try:
            async with create_session_factory(engine)() as session:
                released = await DocumentRepository(session).release_for_retry(uuid.UUID(doc_id))
                await session.commit()
                return released
        finally:
            await engine.dispose()

    assert asyncio.run(release()) is True
    _assert_no_chunks_remain(migrated_database_url, doc_id)
    assert row_of(migrated_database_url, doc_id)["status"] == "pending"


def test_rechunk_task_ignores_invalid_ids(
    auth_client: TestClient,
) -> None:
    from app.workers import document_tasks
    from app.workers.celery_app import create_celery_app

    create_celery_app(auth_client.app.state.settings, configure_logs=False).set_current()  # type: ignore[attr-defined]

    result = document_tasks.rechunk_document.apply(args=["not-a-uuid"])

    assert result.get() == {"status": "skipped", "reason": "invalid_id"}
