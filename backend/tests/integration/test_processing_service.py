import asyncio
import io
import json
import logging
import uuid
import zipfile
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from app.core.documents.processing import SAFE_MESSAGES, FailureReason
from app.observability.logging import JsonFormatter, RequestContextFilter
from app.storage.base import StorageError
from app.storage.local import LocalStorageProvider
from app.workers import document_tasks
from tests.conftest import ClientFactory
from tests.files import make_user, markdown_bytes, text_bytes, upload_ok
from tests.helpers import db_execute, db_scalar
from tests.pdf_factory import make_pdf
from tests.processing_helpers import fetch_document, process_now, row_of, sections_of

pytestmark = pytest.mark.integration

LOCKED_PDF = make_pdf(["secret contents here"], encrypt_password="pw")
BANGLA_LINE = "বাংলাদেশের রাজধানী ঢাকা। এটি একটি পরীক্ষামূলক বাক্য।"


def _docx(*paragraphs: tuple[str, int | None]) -> bytes:
    document = docx.Document()
    for text, level in paragraphs:
        if level:
            document.add_heading(text, level=level)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, headers: dict[str, str], name: str, data: bytes) -> str:
    return str(upload_ok(client, headers, filename=name, content=data, content_type=None)["id"])


# --- successful processing for every format ------------------------------------------------


def test_pdf_is_processed_into_one_section_per_page_with_bangla_intact(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    pdf = make_pdf(
        [f"{BANGLA_LINE}\nPage one in Bangla", "", "Mixed English and বাংলা ক‍্ষ text", "Last page"]
    )
    doc_id = _upload(auth_client, alice, "report.pdf", pdf)

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "chunked"
    sections = sections_of(migrated_database_url, doc_id)
    assert [(s["kind"], s["page_number"], s["ordinal"]) for s in sections] == [
        ("page", 1, 0), ("page", 2, 1), ("page", 3, 2), ("page", 4, 3),
    ]  # fmt: skip
    assert BANGLA_LINE in sections[0]["text"]
    assert sections[1]["text"] == "" and sections[1]["char_count"] == 0  # empty page kept
    assert "ক‍্ষ" in sections[2]["text"]  # ZWJ preserved
    assert all(s["char_count"] == len(s["text"]) for s in sections)


def test_document_row_reports_status_counts_and_metadata(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    pdf = make_pdf(
        [BANGLA_LINE * 3, "Second page"], metadata={"title": "শিরোনাম", "author": "Writer"}
    )
    doc_id = _upload(auth_client, alice, "report.pdf", pdf)

    process_now(auth_client, doc_id)

    row = row_of(migrated_database_url, doc_id)
    assert (
        row["status"] == "chunked"
        and row["failure_reason"] is None
        and row["error_message"] is None
    )
    assert row["page_count"] == 2
    assert row["character_count"] == sum(
        s["char_count"] for s in sections_of(migrated_database_url, doc_id)
    )
    assert row["processing_attempts"] == 1
    assert row["processing_started_at"] <= row["processing_completed_at"]
    meta = row["metadata"]
    assert meta["extractor"] == "pdf" and meta["processing_version"].startswith("p4")
    assert meta["duration_ms"] >= 0 and meta["section_count"] == 2
    assert meta["primary_script"] in {"bengali", "mixed"}
    assert meta["properties"]["title"] == "শিরোনাম" and meta["properties"]["author"] == "Writer"
    assert set(meta["scripts"]) == {"bengali", "latin", "other"}


def test_api_exposes_safe_processing_information(auth_client: TestClient) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "report.pdf", make_pdf(["one", "two"]))
    process_now(auth_client, doc_id)

    data = fetch_document(auth_client, alice, doc_id)

    assert data["status"] == "chunked" and data["page_count"] == 2 and data["character_count"] > 0
    assert data["processing_started_at"] and data["processing_completed_at"]
    assert data["failure_reason"] is None and data["error_message"] is None
    assert "metadata" not in data and "processing_attempts" not in data


def test_docx_is_split_at_headings_and_keeps_bangla(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    data = _docx(
        ("Preamble", None),
        ("ভূমিকা", 1),
        (BANGLA_LINE, None),
        ("Details", 2),
        ("English text", None),
    )
    doc_id = _upload(auth_client, alice, "notes.docx", data)

    assert process_now(auth_client, doc_id).status == "chunked"

    sections = sections_of(migrated_database_url, doc_id)
    assert [(s["kind"], s["heading"], s["heading_level"]) for s in sections] == [
        ("body", None, None), ("section", "ভূমিকা", 1), ("section", "Details", 2),
    ]  # fmt: skip
    assert sections[1]["text"] == BANGLA_LINE
    assert row_of(migrated_database_url, doc_id)["page_count"] is None


def test_text_and_markdown_are_processed(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    txt = _upload(
        auth_client, alice, "plain.txt", f"{BANGLA_LINE}\r\n\r\n\r\n\r\nSecond   para  \n".encode()
    )
    md = _upload(
        auth_client, alice, "guide.md", markdown_bytes("x") + "\n## দ্বিতীয়\nবিষয়বস্তু\n".encode()
    )

    assert process_now(auth_client, txt).status == "chunked"
    assert process_now(auth_client, md).status == "chunked"

    body = sections_of(migrated_database_url, txt)
    assert (
        len(body) == 1 and body[0]["text"] == f"{BANGLA_LINE}\n\nSecond   para"
    )  # blank runs collapsed, spacing kept
    headings = [(s["heading"], s["heading_level"]) for s in sections_of(migrated_database_url, md)]
    assert headings == [("Title x", 1), ("দ্বিতীয়", 2)]


def test_text_is_unicode_normalised_before_storage(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    raw = "Café কো​ x y­ z".encode()
    doc_id = _upload(auth_client, alice, "u.txt", raw)

    process_now(auth_client, doc_id)

    assert sections_of(migrated_database_url, doc_id)[0]["text"] == "Café কো x y z"


# --- failures ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "data", "reason"),
    [
        ("corrupt.pdf", b"%PDF-1.4\nnot really a pdf\n%%EOF\n", FailureReason.CORRUPT_DOCUMENT),
        ("blank.pdf", make_pdf(["", ""]), FailureReason.EMPTY_DOCUMENT),
        ("locked.pdf", LOCKED_PDF, FailureReason.ENCRYPTED_DOCUMENT),
        ("empty.docx", _docx(("", None)), FailureReason.EMPTY_DOCUMENT),
        ("spaces.txt", b"   \n\n \t  \n", FailureReason.EMPTY_DOCUMENT),
    ],
    ids=["corrupt-pdf", "blank-pdf", "encrypted-pdf", "empty-docx", "blank-text"],
)  # fmt: skip
def test_unprocessable_documents_fail_with_a_safe_reason(
    auth_client: TestClient,
    migrated_database_url: str,
    name: str,
    data: bytes,
    reason: FailureReason,
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, name, data)

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed" and outcome.reason == reason.value
    shown = fetch_document(auth_client, alice, doc_id)
    assert shown["status"] == "failed" and shown["failure_reason"] == reason.value
    assert shown["error_message"] == SAFE_MESSAGES[reason]
    assert shown["processing_completed_at"]
    assert sections_of(migrated_database_url, doc_id) == []


def test_corrupt_docx_with_valid_container_but_bad_xml_fails(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document><oops>")
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "bad.docx", buffer.getvalue())

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed" and outcome.reason == "corrupt_document"


def test_missing_storage_object_fails_without_leaking_paths(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "gone.txt", text_bytes("gone"))
    root = Path(auth_client.app.state.settings.storage_local_path)  # type: ignore[attr-defined]
    for path in root.rglob("*"):
        if path.is_file():
            path.unlink()

    outcome = process_now(auth_client, doc_id)

    assert outcome.reason == "storage_missing"
    shown = fetch_document(auth_client, alice, doc_id)
    assert str(root) not in json.dumps(shown) and "documents/" not in json.dumps(shown)


def test_file_whose_size_changed_on_disk_is_rejected(auth_client: TestClient) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("size check"))
    root = Path(auth_client.app.state.settings.storage_local_path)  # type: ignore[attr-defined]
    for path in root.rglob("*"):
        if path.is_file():
            path.write_bytes(path.read_bytes() + b"extra")

    assert process_now(auth_client, doc_id).reason == "corrupt_document"


def test_page_and_text_limits_are_configurable(auth_client_factory: ClientFactory) -> None:
    client = auth_client_factory(processing_max_pages=2, processing_max_text_chars=1000)
    alice = make_user(client, "alice@example.com")
    many_pages = _upload(client, alice, "long.pdf", make_pdf(["a b c", "d e f", "g h i"]))
    big_text = _upload(client, alice, "big.txt", b"word " * 400)

    assert process_now(client, many_pages).reason == "too_many_pages"
    assert process_now(client, big_text).reason == "content_too_large"


def test_processing_time_budget_is_enforced(
    auth_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.core.documents.processing import Deadline
    from app.services import processing_service

    monkeypatch.setattr(processing_service, "Deadline", lambda seconds: Deadline(-1))
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.pdf", make_pdf(["some text", "more text"]))

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed" and outcome.reason == "timeout"


def test_unexpected_parser_error_is_contained_and_not_retried(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.parsers.pdf_parser import PDFExtractor

    def explode(self: PDFExtractor, *args: object, **kwargs: object) -> None:
        raise RuntimeError("boom with secret internals /srv/app/x.py")

    monkeypatch.setattr(PDFExtractor, "extract", explode)
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.pdf", make_pdf(["text on page"]))

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "failed" and outcome.reason == "extraction_failed"
    shown = json.dumps(fetch_document(auth_client, alice, doc_id))
    assert "boom" not in shown and "secret internals" not in shown and "/srv/app" not in shown
    assert row_of(migrated_database_url, doc_id)["processing_attempts"] == 1


# --- retries, idempotency, concurrency ---------------------------------------------------


def test_transient_storage_error_is_retried_and_then_succeeds(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("retry me"))
    real_size = LocalStorageProvider.size
    state = {"calls": 0}

    async def flaky_size(self: LocalStorageProvider, key: str) -> int:
        state["calls"] += 1
        if state["calls"] == 1:
            raise StorageError("temporarily unavailable")
        return await real_size(self, key)

    monkeypatch.setattr(LocalStorageProvider, "size", flaky_size)

    first = process_now(auth_client, doc_id)
    after_first = row_of(migrated_database_url, doc_id)
    second = process_now(auth_client, doc_id)

    backoff = auth_client.app.state.settings.processing_retry_backoff_seconds  # type: ignore[attr-defined]
    assert first.status == "retry" and first.retry_in == backoff
    assert after_first["status"] == "pending" and after_first["processing_attempts"] == 1
    assert second.status == "chunked"
    final = row_of(migrated_database_url, doc_id)
    assert final["status"] == "chunked" and final["processing_attempts"] == 2
    assert len(sections_of(migrated_database_url, doc_id)) == 1  # no duplicates


def test_retry_budget_is_bounded_and_ends_in_failed(
    auth_client_factory: ClientFactory, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    client = auth_client_factory(processing_max_attempts=2)
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "t.txt", text_bytes("always failing"))

    async def always_down(self: LocalStorageProvider, key: str) -> int:
        raise StorageError("down")

    monkeypatch.setattr(LocalStorageProvider, "size", always_down)

    outcomes = [process_now(client, doc_id).status for _ in range(3)]

    assert outcomes == ["retry", "failed", "skipped"]
    row = row_of(migrated_database_url, doc_id)
    assert row["status"] == "failed" and row["failure_reason"] == "storage_unavailable"
    assert row["processing_attempts"] == 2


def test_backoff_grows_exponentially_and_is_capped(auth_client_factory: ClientFactory) -> None:
    from app.services.processing_service import MAX_BACKOFF_SECONDS, ProcessingService

    client = auth_client_factory(processing_retry_backoff_seconds=30)
    service = ProcessingService(
        session=None,  # type: ignore[arg-type]
        documents=None,  # type: ignore[arg-type]
        sections=None,  # type: ignore[arg-type]
        chunks=None,  # type: ignore[arg-type]
        chunker=None,  # type: ignore[arg-type]
        storage=None,  # type: ignore[arg-type]
        settings=client.app.state.settings,  # type: ignore[attr-defined]
    )

    assert [service._backoff(n) for n in (1, 2, 3, 4)] == [30, 60, 120, 240]
    assert service._backoff(20) == MAX_BACKOFF_SECONDS


def test_only_pending_documents_can_be_claimed(auth_client: TestClient) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("once"))

    first = process_now(auth_client, doc_id)
    again = process_now(auth_client, doc_id)  # already chunked
    unknown = process_now(auth_client, str(uuid.uuid4()))

    assert (first.status, again.status, unknown.status) == ("chunked", "skipped", "skipped")


def test_two_workers_racing_for_one_document_process_it_exactly_once(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.pdf", make_pdf(["racing text", "second page"]))
    settings = auth_client.app.state.settings  # type: ignore[attr-defined]

    async def race() -> list[str]:
        results = await asyncio.gather(
            *(document_tasks._process(settings, uuid.UUID(doc_id), f"w{i}") for i in range(4))
        )
        return [r.status for r in results]

    statuses = asyncio.run(race())

    assert sorted(statuses) == ["chunked", "skipped", "skipped", "skipped"]
    assert row_of(migrated_database_url, doc_id)["processing_attempts"] == 1
    assert len(sections_of(migrated_database_url, doc_id)) == 2


def test_manual_retry_reprocesses_without_duplicating_content(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.pdf", make_pdf(["alpha", "beta"]))
    process_now(auth_client, doc_id)
    db_execute(
        migrated_database_url, "UPDATE documents SET status = 'pending' WHERE id = :d", d=doc_id
    )

    assert process_now(auth_client, doc_id).status == "chunked"

    assert len(sections_of(migrated_database_url, doc_id)) == 2
    assert row_of(migrated_database_url, doc_id)["processing_attempts"] == 2


def test_abandoned_parsing_document_is_reclaimed_after_the_stale_window(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("crashed worker"))
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='parsing', processing_attempts=1, "
        "processing_started_at = now() - interval '1 hour' WHERE id = :d",
        d=doc_id,
    )

    assert process_now(auth_client, doc_id).status == "chunked"


def test_fresh_parsing_document_is_not_stolen_from_a_live_worker(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("in progress"))
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='parsing', processing_attempts=1, "
        "processing_started_at = now() WHERE id = :d",
        d=doc_id,
    )

    assert process_now(auth_client, doc_id).status == "skipped"
    assert row_of(migrated_database_url, doc_id)["status"] == "parsing"


def test_document_deleted_while_processing_leaves_no_orphans(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.services.processing_service import ProcessingService

    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("deleted mid-flight"))
    real_read = ProcessingService._read

    async def read_then_delete(self: ProcessingService, claimed: object) -> bytes:
        data = await real_read(self, claimed)  # type: ignore[arg-type]
        await asyncio.to_thread(
            db_execute, migrated_database_url, "DELETE FROM documents WHERE id = :d", d=doc_id
        )
        return data

    monkeypatch.setattr(ProcessingService, "_read", read_then_delete)

    outcome = process_now(auth_client, doc_id)

    assert outcome.status == "skipped"
    assert db_scalar(migrated_database_url, "SELECT count(*) FROM document_sections") == 0


# --- recovery sweep ----------------------------------------------------------------------


def _recover(client: TestClient) -> tuple[object, list[uuid.UUID]]:
    settings = client.app.state.settings  # type: ignore[attr-defined]
    queued: list[uuid.UUID] = []
    report = asyncio.run(document_tasks._recover(settings, queued.append))
    return report, queued


def test_sweep_releases_stalled_documents_and_requeues_them(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("stalled"))
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='parsing', processing_attempts=1, "
        "processing_started_at = now() - interval '1 hour' WHERE id = :d",
        d=doc_id,
    )

    report, queued = _recover(auth_client)

    assert report == {"released": 1, "abandoned": 0, "requeued": 1}
    assert queued == [uuid.UUID(doc_id)]
    assert row_of(migrated_database_url, doc_id)["status"] == "pending"


def test_sweep_abandons_documents_that_keep_crashing_workers(
    auth_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = auth_client_factory(processing_max_attempts=2)
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "t.txt", text_bytes("poison document"))
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='parsing', processing_attempts=2, "
        "processing_started_at = now() - interval '1 hour' WHERE id = :d",
        d=doc_id,
    )

    report, queued = _recover(client)

    assert report == {"released": 0, "abandoned": 1, "requeued": 0} and queued == []
    row = row_of(migrated_database_url, doc_id)
    assert row["status"] == "failed" and row["failure_reason"] == "retries_exhausted"


def test_sweep_requeues_pending_documents_whose_message_was_lost_only_once_per_window(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("lost message"))
    recent = _recover(auth_client)[1]
    db_execute(
        migrated_database_url,
        "UPDATE documents SET updated_at = now() - interval '1 hour' WHERE id = :d",
        d=doc_id,
    )

    first = _recover(auth_client)[1]
    second = _recover(auth_client)[1]

    assert recent == [] and first == [uuid.UUID(doc_id)] and second == []


def test_sweep_leaves_healthy_documents_alone(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("healthy"))
    process_now(auth_client, doc_id)

    report, queued = _recover(auth_client)

    assert report == {"released": 0, "abandoned": 0, "requeued": 0} and queued == []


# --- logging ----------------------------------------------------------------------------


def test_processing_logs_metadata_but_never_document_content(
    auth_client_factory: ClientFactory,
) -> None:
    lines: list[str] = []

    class Capture(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            lines.append(self.format(record))

    client = auth_client_factory(log_level="DEBUG")
    handler = Capture(level=logging.DEBUG)
    handler.setFormatter(JsonFormatter())
    handler.addFilter(RequestContextFilter("test"))
    logging.getLogger().addHandler(handler)
    try:
        alice = make_user(client, "alice@example.com")
        secret = "TOPSECRET-CONTENT-ওপেন"
        ok = _upload(client, alice, "ok.pdf", make_pdf([f"{secret} page one", "page two"]))
        bad = _upload(
            client, alice, "bad.pdf", b"%PDF-1.4\nbroken " + secret.encode() + b"\n%%EOF\n"
        )
        process_now(client, ok, task_id="task-ok")
        process_now(client, bad, task_id="task-bad")
    finally:
        logging.getLogger().removeHandler(handler)

    logged = "\n".join(lines)
    assert "TOPSECRET" not in logged and "ওপেন" not in logged
    events = [json.loads(line) for line in lines if '"app.processing"' in line]
    extracted = next(e for e in events if e["message"] == "text extracted")
    assert extracted["document_id"] == ok and extracted["task_id"] == "task-ok"
    assert extracted["extractor"] == "pdf" and extracted["page_count"] == 2
    assert extracted["character_count"] > 0
    completed = next(e for e in events if e["message"] == "processing completed")
    assert completed["document_id"] == ok and completed["chunk_count"] >= 2
    assert completed["chunking_version"].startswith("c") and completed["duration_ms"] >= 0
    failed = next(e for e in events if e["message"] == "processing failed")
    assert failed["failure_reason"] == "corrupt_document" and failed["task_id"] == "task-bad"


def test_a_document_that_has_used_up_its_attempts_fails_instead_of_looping(
    auth_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = auth_client_factory(processing_max_attempts=3)
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "t.txt", text_bytes("crash loop"))
    db_execute(
        migrated_database_url,
        "UPDATE documents SET processing_attempts = 3 WHERE id = :d",
        d=doc_id,
    )

    outcome = process_now(client, doc_id)

    assert outcome.status == "failed" and outcome.reason == "retries_exhausted"
    assert row_of(migrated_database_url, doc_id)["failure_reason"] == "retries_exhausted"


def test_sweep_still_releases_documents_when_the_broker_rejects_the_requeue(
    auth_client: TestClient, migrated_database_url: str
) -> None:
    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("broker down"))
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='parsing', processing_attempts=1, "
        "processing_started_at = now() - interval '1 hour' WHERE id = :d",
        d=doc_id,
    )

    def broken_enqueue(document_id: uuid.UUID) -> None:
        raise ConnectionError("broker down")

    report = asyncio.run(
        document_tasks._recover(auth_client.app.state.settings, broken_enqueue)  # type: ignore[attr-defined]
    )

    assert report == {"released": 1, "abandoned": 0, "requeued": 0}
    assert row_of(migrated_database_url, doc_id)["status"] == "pending"  # the next sweep retries


def test_soft_time_limit_marks_the_document_failed_with_a_timeout(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from celery.exceptions import SoftTimeLimitExceeded

    from app.workers.celery_app import create_celery_app

    alice = make_user(auth_client, "alice@example.com")
    doc_id = _upload(auth_client, alice, "t.txt", text_bytes("too slow"))
    db_execute(
        migrated_database_url, "UPDATE documents SET status='parsing' WHERE id = :d", d=doc_id
    )

    async def interrupted(*args: object, **kwargs: object) -> None:
        raise SoftTimeLimitExceeded

    monkeypatch.setattr(document_tasks, "_process", interrupted)
    create_celery_app(auth_client.app.state.settings, configure_logs=False).set_current()  # type: ignore[attr-defined]

    result = document_tasks.process_document.apply(args=[doc_id])

    assert result.get() == {"status": "failed", "reason": "timeout"}
    row = row_of(migrated_database_url, doc_id)
    assert row["status"] == "failed" and row["failure_reason"] == "timeout"
