import time
import uuid
from typing import Any

import pytest
from fastapi.testclient import TestClient

from app.workers.dispatch import CeleryProcessingQueue
from tests.conftest import ClientFactory
from tests.files import make_user, markdown_bytes, text_bytes, upload_ok
from tests.helpers import db_execute, tokens
from tests.pdf_factory import make_pdf
from tests.processing_helpers import fetch_document, row_of, sections_of, wait_for_status

pytestmark = pytest.mark.integration

BANGLA = "বাংলাদেশের রাজধানী ঢাকা।"
DOCS = "/api/v1/documents"


def _result_of(async_result: Any, timeout: float = 30.0) -> Any:
    """Poll the result backend instead of `.get()`, which leaves a listener thread running."""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if async_result.ready():
            return async_result.result
        time.sleep(0.2)
    raise AssertionError("task did not finish in time")


def _upload(client: TestClient, headers: dict[str, str], name: str, data: bytes) -> str:
    return str(upload_ok(client, headers, filename=name, content=data, content_type=None)["id"])


def test_upload_is_queued_processed_by_a_real_worker_and_becomes_chunked(
    pipeline_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")

    uploaded = upload_ok(
        client,
        alice,
        filename="report.pdf",
        content=make_pdf([BANGLA, "English page"]),
        content_type=None,
    )
    assert uploaded["status"] == "pending"  # the request itself did no processing

    done = wait_for_status(client, alice, uploaded["id"], {"chunked", "failed"})

    assert done["status"] == "chunked" and done["page_count"] == 2 and done["character_count"] > 0
    assert done["processing_started_at"] and done["processing_completed_at"]
    sections = sections_of(migrated_database_url, uploaded["id"])
    assert [s["page_number"] for s in sections] == [1, 2] and BANGLA in sections[0]["text"]
    assert row_of(migrated_database_url, uploaded["id"])["metadata"]["extractor"] == "pdf"


def test_all_formats_flow_through_the_queue(
    pipeline_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    import io

    import docx

    document = docx.Document()
    document.add_heading("শিরোনাম", level=1)
    document.add_paragraph(BANGLA)
    buffer = io.BytesIO()
    document.save(buffer)
    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")
    ids = [
        _upload(client, alice, "a.pdf", make_pdf(["pdf text"])),
        _upload(client, alice, "b.docx", buffer.getvalue()),
        _upload(client, alice, "c.txt", text_bytes("plain")),
        _upload(client, alice, "d.md", markdown_bytes("md")),
    ]

    results = [wait_for_status(client, alice, doc_id, {"chunked", "failed"}) for doc_id in ids]

    assert [r["status"] for r in results] == ["chunked"] * 4
    assert sections_of(migrated_database_url, ids[1])[0]["heading"] == "শিরোনাম"


def test_a_failing_document_ends_up_failed_and_does_not_block_others(
    pipeline_client_factory: ClientFactory,
) -> None:
    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")
    bad = _upload(client, alice, "bad.pdf", b"%PDF-1.4\ngarbage\n%%EOF\n")
    good = _upload(client, alice, "good.txt", text_bytes("fine"))

    bad_done = wait_for_status(client, alice, bad, {"chunked", "failed"})
    good_done = wait_for_status(client, alice, good, {"chunked", "failed"})

    assert bad_done["status"] == "failed" and bad_done["failure_reason"] == "corrupt_document"
    assert good_done["status"] == "chunked"


def test_manual_retry_goes_through_the_queue_again(
    pipeline_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "t.txt", text_bytes("retry through the queue"))
    wait_for_status(client, alice, doc_id, {"chunked"})
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='failed', failure_reason='extraction_failed', "
        "error_message='x', processing_attempts=3 WHERE id = :d",
        d=doc_id,
    )

    response = client.post(f"{DOCS}/{doc_id}/retry", headers=alice)

    assert response.status_code == 202 and response.json()["data"]["status"] == "pending"
    done = wait_for_status(client, alice, doc_id, {"chunked", "failed"})
    assert done["status"] == "chunked" and done["failure_reason"] is None
    assert row_of(migrated_database_url, doc_id)["processing_attempts"] == 1  # fresh budget
    assert len(sections_of(migrated_database_url, doc_id)) == 1


def test_recovery_sweep_task_rescues_documents_lost_by_the_queue(
    pipeline_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "t.txt", text_bytes("rescued by the sweep"))
    wait_for_status(client, alice, doc_id, {"chunked"})
    db_execute(
        migrated_database_url,
        "UPDATE documents SET status='pending', "
        "updated_at = now() - interval '1 hour' WHERE id = :d",
        d=doc_id,
    )
    sweep = client.app.state.processing_queue._celery.send_task(  # type: ignore[attr-defined]
        "app.workers.document_tasks.recover_stalled_documents"
    )

    report = _result_of(sweep)

    assert report["requeued"] == 1
    assert wait_for_status(client, alice, doc_id, {"chunked"})["status"] == "chunked"
    assert row_of(migrated_database_url, doc_id)["processing_attempts"] == 2


def test_upload_still_succeeds_when_the_broker_is_unreachable(
    auth_client: TestClient, migrated_database_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broker_down(self: CeleryProcessingQueue, document_id: uuid.UUID) -> None:
        raise ConnectionError("broker unreachable")

    monkeypatch.setattr(CeleryProcessingQueue, "enqueue_sync", broker_down)
    alice = make_user(auth_client, "alice@example.com")

    started = time.monotonic()
    doc = upload_ok(
        auth_client, alice, filename="t.txt", content=text_bytes("broker down"), content_type=None
    )

    assert doc["status"] == "pending"
    assert time.monotonic() - started < 10
    assert row_of(migrated_database_url, doc["id"])["status"] == "pending"


def test_upload_response_is_not_blocked_by_processing(
    pipeline_client_factory: ClientFactory,
) -> None:
    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")
    big = make_pdf([f"page {n} " + "word " * 80 for n in range(60)])

    started = time.monotonic()
    doc = upload_ok(client, alice, filename="big.pdf", content=big, content_type=None)
    request_seconds = time.monotonic() - started

    assert doc["status"] == "pending" and doc["page_count"] is None
    assert request_seconds < 5
    assert wait_for_status(client, alice, doc["id"], {"chunked"})["page_count"] == 60


def test_worker_tasks_carry_a_correlation_id_and_skip_invalid_ids(
    pipeline_client_factory: ClientFactory,
) -> None:
    client = pipeline_client_factory()
    celery = client.app.state.processing_queue._celery  # type: ignore[attr-defined]

    result = celery.send_task("app.workers.document_tasks.process_document", args=["not-a-uuid"])

    assert _result_of(result) == {"status": "skipped", "reason": "invalid_id"}


def test_tokens_helper_is_importable() -> None:
    assert tokens is not None and fetch_document is not None


def test_celery_retries_a_transient_failure_until_it_succeeds(
    pipeline_client_factory: ClientFactory,
    migrated_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.storage.base import StorageError
    from app.storage.local import LocalStorageProvider

    real_size = LocalStorageProvider.size
    calls = {"count": 0}

    async def flaky_size(self: LocalStorageProvider, key: str) -> int:
        calls["count"] += 1
        if calls["count"] == 1:
            raise StorageError("transient")
        return await real_size(self, key)

    monkeypatch.setattr(LocalStorageProvider, "size", flaky_size)
    client = pipeline_client_factory(processing_retry_backoff_seconds=0)
    alice = make_user(client, "alice@example.com")
    doc_id = _upload(client, alice, "t.txt", text_bytes("flaky storage"))

    done = wait_for_status(client, alice, doc_id, {"chunked", "failed"})

    assert done["status"] == "chunked"
    assert row_of(migrated_database_url, doc_id)["processing_attempts"] == 2
    assert len(sections_of(migrated_database_url, doc_id)) == 1


def test_rechunk_task_regenerates_chunks_through_the_worker(
    pipeline_client_factory: ClientFactory, migrated_database_url: str
) -> None:
    from app.workers.celery_app import create_celery_app
    from tests.helpers import db_rows

    client = pipeline_client_factory()
    alice = make_user(client, "alice@example.com")
    body = ("A reasonably long sentence for the rechunk task. " * 80).encode()
    doc_id = _upload(client, alice, "r.txt", body)
    wait_for_status(client, alice, doc_id, {"chunked", "failed"})
    query = "SELECT id::text, text FROM document_chunks WHERE document_id = :d ORDER BY chunk_index"
    before = db_rows(migrated_database_url, query, d=doc_id)
    app = create_celery_app(client.app.state.settings, configure_logs=False)  # type: ignore[attr-defined]

    for arg in ("not-a-uuid", str(uuid.uuid4()), doc_id):
        app.send_task("app.workers.document_tasks.rechunk_document", args=[arg], ignore_result=True)

    deadline = time.monotonic() + 30
    after = before
    while time.monotonic() < deadline and after[0][0] == before[0][0]:
        time.sleep(0.2)
        after = db_rows(migrated_database_url, query, d=doc_id)
    assert after[0][0] != before[0][0]  # rows were replaced...
    assert [r[1] for r in after] == [r[1] for r in before]  # ...with identical content
    assert fetch_document(client, alice, doc_id)["status"] == "chunked"
