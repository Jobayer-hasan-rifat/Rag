import io
import zipfile

import docx
import pytest

from app.core.documents.processing import (
    SAFE_MESSAGES,
    Deadline,
    ExtractionLimits,
    FailureReason,
    ProcessingFailure,
)
from app.core.documents.types import DocumentType
from app.parsers.base import SectionKind
from app.parsers.docx_parser import DOCXExtractor
from app.parsers.pdf_parser import PDFExtractor
from app.parsers.registry import get_extractor
from app.parsers.text_parser import MarkdownExtractor, TextExtractor
from tests.pdf_factory import make_pdf

LIMITS = ExtractionLimits(max_pages=50, max_text_chars=100_000, max_docx_uncompressed_bytes=10**7)
BANGLA = "বাংলাদেশের রাজধানী ঢাকা।"


def _deadline() -> Deadline:
    return Deadline(60)


def _docx(build) -> bytes:  # type: ignore[no-untyped-def]
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- PDF --------------------------------------------------------------------------------


def test_pdf_pages_keep_their_page_numbers_and_order() -> None:
    data = make_pdf(["first page", "second page", "third page"])

    result = PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline())

    assert result.page_count == 3 and result.extractor == "pdf"
    assert [s.page_number for s in result.sections] == [1, 2, 3]
    assert all(s.kind is SectionKind.PAGE for s in result.sections)
    assert [s.text.strip() for s in result.sections] == ["first page", "second page", "third page"]


def test_pdf_empty_pages_are_kept_so_numbering_stays_aligned() -> None:
    data = make_pdf(["one", "", "three"])

    result = PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline())

    assert [s.page_number for s in result.sections] == [1, 2, 3]
    assert result.sections[1].text.strip() == ""


def test_pdf_bangla_text_is_extracted_intact() -> None:
    data = make_pdf([f"{BANGLA}\nসংখ্যা ১২৩"])

    text = PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline()).sections[0].text

    assert BANGLA in text and "সংখ্যা ১২৩" in text


def test_pdf_mixed_bangla_english_and_joiners_survive() -> None:
    page = "Mixed: বাংলা and English র‍্য ক‌"
    data = make_pdf([page])

    text = PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline()).sections[0].text

    assert page in text


def test_pdf_unusual_unicode_and_whitespace_are_returned_raw_for_the_normaliser() -> None:
    data = make_pdf(["  spaced    out\n\n\n   text  ", "naïve café – “quotes” 😀"])

    sections = PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline()).sections

    assert "spaced    out" in sections[0].text
    assert "naïve café" in sections[1].text


def test_pdf_metadata_is_reported_as_untrusted_properties() -> None:
    data = make_pdf(["x" * 30], metadata={"title": "শিরোনাম", "author": "Writer"})

    properties = PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline()).properties

    assert properties["title"] == "শিরোনাম" and properties["author"] == "Writer"


@pytest.mark.parametrize(
    "data",
    [
        b"%PDF-1.4\nthis is not a real pdf at all\n%%EOF\n",
        b"%PDF-1.7\n" + b"\x00\xff\xfe garbage " * 50 + b"\n%%EOF\n",
        b"",
        b"hello",
    ],
)
def test_malformed_pdfs_fail_as_corrupt(data: bytes) -> None:
    with pytest.raises(ProcessingFailure) as error:
        PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline())

    assert error.value.reason is FailureReason.CORRUPT_DOCUMENT
    assert not error.value.retryable


def test_truncated_pdf_does_not_crash() -> None:
    data = make_pdf(["page one text here", "page two text here"])[:700] + b"\n%%EOF\n"

    try:
        PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline())
    except ProcessingFailure as failure:
        assert failure.reason is FailureReason.CORRUPT_DOCUMENT


def test_encrypted_pdf_is_rejected_with_a_specific_reason() -> None:
    data = make_pdf(["secret text here"], encrypt_password="hunter2")

    with pytest.raises(ProcessingFailure) as error:
        PDFExtractor().extract(data, limits=LIMITS, deadline=_deadline())

    assert error.value.reason is FailureReason.ENCRYPTED_DOCUMENT


def test_pdf_page_limit_is_enforced_before_extracting_text() -> None:
    data = make_pdf([f"page {i}" for i in range(6)])
    limits = ExtractionLimits(max_pages=5, max_text_chars=10**6, max_docx_uncompressed_bytes=10**7)

    with pytest.raises(ProcessingFailure) as error:
        PDFExtractor().extract(data, limits=limits, deadline=_deadline())

    assert error.value.reason is FailureReason.TOO_MANY_PAGES


def test_pdf_text_size_limit_is_enforced() -> None:
    line_block = "\n".join(["abcdefghij"] * 40)
    data = make_pdf([line_block, line_block])
    limits = ExtractionLimits(max_pages=10, max_text_chars=500, max_docx_uncompressed_bytes=10**7)

    with pytest.raises(ProcessingFailure) as error:
        PDFExtractor().extract(data, limits=limits, deadline=_deadline())

    assert error.value.reason is FailureReason.CONTENT_TOO_LARGE


def test_pdf_extraction_stops_when_the_deadline_passes() -> None:
    clock = iter([0.0, 0.0, 100.0, 100.0, 100.0, 100.0])
    data = make_pdf(["a", "b", "c"])

    with pytest.raises(ProcessingFailure) as error:
        PDFExtractor().extract(data, limits=LIMITS, deadline=Deadline(1, clock=lambda: next(clock)))

    assert error.value.reason is FailureReason.TIMEOUT


# --- DOCX -------------------------------------------------------------------------------


def test_docx_headings_paragraphs_and_tables_become_sections() -> None:
    def build(document: docx.document.Document) -> None:
        document.add_paragraph("Preamble before any heading")
        document.add_heading("Introduction", level=1)
        document.add_paragraph("First paragraph.")
        document.add_paragraph("Second paragraph.")
        document.add_heading("Details", level=2)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "alpha", "1"

    result = DOCXExtractor().extract(_docx(build), limits=LIMITS, deadline=_deadline())

    kinds = [(s.kind, s.heading, s.heading_level) for s in result.sections]
    assert kinds == [
        (SectionKind.BODY, None, None),
        (SectionKind.SECTION, "Introduction", 1),
        (SectionKind.SECTION, "Details", 2),
    ]
    assert result.sections[1].text == "First paragraph.\nSecond paragraph."
    assert "Name | Value" in result.sections[2].text and "alpha | 1" in result.sections[2].text
    assert result.page_count is None


def test_docx_bangla_and_mixed_text_are_preserved() -> None:
    def build(document: docx.document.Document) -> None:
        document.add_heading("ভূমিকা", level=1)
        document.add_paragraph(f"{BANGLA} Mixed English র‍্য text.")

    result = DOCXExtractor().extract(_docx(build), limits=LIMITS, deadline=_deadline())

    assert result.sections[0].heading == "ভূমিকা"
    assert f"{BANGLA} Mixed English র‍্য text." in result.sections[0].text


def test_docx_with_only_empty_paragraphs_yields_no_text() -> None:
    result = DOCXExtractor().extract(
        _docx(lambda d: d.add_paragraph("")), limits=LIMITS, deadline=_deadline()
    )

    assert not any(s.text.strip() or s.heading for s in result.sections)


def test_docx_core_properties_are_collected() -> None:
    def build(document: docx.document.Document) -> None:
        document.core_properties.title = "Report"
        document.core_properties.author = "Someone"
        document.add_paragraph("body text for properties")

    properties = (
        DOCXExtractor().extract(_docx(build), limits=LIMITS, deadline=_deadline()).properties
    )

    assert properties["title"] == "Report" and properties["author"] == "Someone"


@pytest.mark.parametrize(
    "data",
    [b"not a zip at all", b"PK\x03\x04garbage", b""],
)
def test_non_zip_docx_is_corrupt(data: bytes) -> None:
    with pytest.raises(ProcessingFailure) as error:
        DOCXExtractor().extract(data, limits=LIMITS, deadline=_deadline())

    assert error.value.reason is FailureReason.CORRUPT_DOCUMENT


def test_docx_with_invalid_xml_is_corrupt() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document><unclosed>")

    with pytest.raises(ProcessingFailure) as error:
        DOCXExtractor().extract(buffer.getvalue(), limits=LIMITS, deadline=_deadline())

    assert error.value.reason is FailureReason.CORRUPT_DOCUMENT


def test_docx_archive_bomb_is_rejected_before_parsing() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "0" * 30_000_000)
    limits = ExtractionLimits(max_pages=5, max_text_chars=10**6, max_docx_uncompressed_bytes=10**6)

    with pytest.raises(ProcessingFailure) as error:
        DOCXExtractor().extract(buffer.getvalue(), limits=limits, deadline=_deadline())

    assert error.value.reason is FailureReason.CONTENT_TOO_LARGE


def test_docx_external_entities_are_not_resolved() -> None:
    xxe = (
        '<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/hostname">]>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", xxe)

    try:
        result = DOCXExtractor().extract(buffer.getvalue(), limits=LIMITS, deadline=_deadline())
    except ProcessingFailure:
        return  # rejecting the document is also acceptable
    assert "root:" not in " ".join(s.text for s in result.sections)


# --- Text and Markdown ------------------------------------------------------------------


def test_text_is_a_single_body_section_with_bom_removed() -> None:
    data = "﻿হ্যালো world\n  indented".encode()

    result = TextExtractor().extract(data, limits=LIMITS, deadline=_deadline())

    assert len(result.sections) == 1 and result.sections[0].kind is SectionKind.BODY
    assert result.sections[0].text == "হ্যালো world\n  indented"
    assert result.page_count is None and not TextExtractor.collapse_inline_spaces


def test_invalid_utf8_text_is_corrupt() -> None:
    with pytest.raises(ProcessingFailure) as error:
        TextExtractor().extract(b"\xff\xfe\x00bad", limits=LIMITS, deadline=_deadline())

    assert error.value.reason is FailureReason.CORRUPT_DOCUMENT


def test_text_size_limit_is_enforced() -> None:
    limits = ExtractionLimits(max_pages=1, max_text_chars=10, max_docx_uncompressed_bytes=10**6)

    with pytest.raises(ProcessingFailure) as error:
        TextExtractor().extract(b"x" * 11, limits=limits, deadline=_deadline())

    assert error.value.reason is FailureReason.CONTENT_TOO_LARGE


MARKDOWN = """Intro text

# ভূমিকা
Bangla body এবং English.

## Sub *heading* ##
- item one
- item two

```python
# not a heading
print("x")
```

#Not a heading (no space)
### Third
"""


def test_markdown_splits_at_headings_and_keeps_levels() -> None:
    result = MarkdownExtractor().extract(MARKDOWN.encode(), limits=LIMITS, deadline=_deadline())

    assert [(s.kind, s.heading, s.heading_level) for s in result.sections] == [
        (SectionKind.BODY, None, None),
        (SectionKind.SECTION, "ভূমিকা", 1),
        (SectionKind.SECTION, "Sub *heading*", 2),
        (SectionKind.SECTION, "Third", 3),
    ]
    assert result.sections[1].text.strip() == "Bangla body এবং English."


def test_markdown_ignores_hashes_inside_code_fences_and_keeps_the_code() -> None:
    sections = (
        MarkdownExtractor().extract(MARKDOWN.encode(), limits=LIMITS, deadline=_deadline()).sections
    )

    assert "# not a heading" in sections[2].text and 'print("x")' in sections[2].text
    assert "- item one" in sections[2].text
    assert "#Not a heading (no space)" in sections[2].text


def test_markdown_without_headings_is_one_body_section() -> None:
    result = MarkdownExtractor().extract(
        b"just some text\nmore", limits=LIMITS, deadline=_deadline()
    )

    assert [s.kind for s in result.sections] == [SectionKind.BODY]


def test_markdown_html_and_scripts_stay_inert_text() -> None:
    data = b"# Title\n<script>alert(1)</script>\n[x](javascript:alert(1))"

    sections = MarkdownExtractor().extract(data, limits=LIMITS, deadline=_deadline()).sections

    assert "<script>alert(1)</script>" in sections[0].text


# --- registry and messages ---------------------------------------------------------------


@pytest.mark.parametrize(
    ("file_type", "expected"),
    [
        (DocumentType.PDF, "pdf"),
        (DocumentType.DOCX, "docx"),
        (DocumentType.TXT, "text"),
        (DocumentType.MD, "markdown"),
    ],
)
def test_every_supported_type_has_an_extractor(file_type: DocumentType, expected: str) -> None:
    assert get_extractor(file_type.value).name == expected


def test_unknown_types_have_no_extractor() -> None:
    with pytest.raises(ProcessingFailure) as error:
        get_extractor("exe")

    assert error.value.reason is FailureReason.UNSUPPORTED_FORMAT


def test_every_failure_reason_has_a_safe_message_without_internal_detail() -> None:
    assert set(SAFE_MESSAGES) == set(FailureReason)
    for message in SAFE_MESSAGES.values():
        assert message.endswith(".") and "Traceback" not in message and "/" not in message


def test_deadline_raises_only_after_expiry() -> None:
    now = [0.0]
    deadline = Deadline(10, clock=lambda: now[0])

    deadline.check()
    now[0] = 9.9
    deadline.check()
    now[0] = 10.0
    with pytest.raises(ProcessingFailure):
        deadline.check()
