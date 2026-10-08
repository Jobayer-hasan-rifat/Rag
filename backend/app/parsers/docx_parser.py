import io
import re
import zipfile
from typing import Any

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.documents.processing import (
    Deadline,
    ExtractionLimits,
    FailureReason,
    ProcessingFailure,
)
from app.parsers.base import DocumentExtractor, ExtractedSection, ExtractionResult, SectionKind

_MAX_ZIP_ENTRIES = 5000
_MAX_COMPRESSION_RATIO = 200
_HEADING_STYLE = re.compile(r"^Heading\s+([1-9])$", re.IGNORECASE)
_DEADLINE_EVERY = 200


class DOCXExtractor(DocumentExtractor):
    """Paragraph, heading and table text from the document body via python-docx.

    The archive is never extracted to disk; macros, embedded objects, headers/footers and
    comments are ignored. python-docx parses XML without resolving external entities.
    """

    name = "docx"

    def extract(
        self, data: bytes, *, limits: ExtractionLimits, deadline: Deadline
    ) -> ExtractionResult:
        _check_archive(data, limits)
        try:
            document = docx.Document(io.BytesIO(data))
            body = list(document.element.body.iterchildren())
        except Exception as error:
            raise ProcessingFailure(
                FailureReason.CORRUPT_DOCUMENT, detail=type(error).__name__
            ) from None

        builder = _SectionBuilder(limits, _heading_levels(document))
        try:
            for index, child in enumerate(body):
                if index % _DEADLINE_EVERY == 0:
                    deadline.check()
                tag = child.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    builder.add_paragraph(Paragraph(child, document))
                elif tag == "tbl":
                    builder.add_table(Table(child, document))
            properties = _properties(document)
        except ProcessingFailure:
            raise
        except Exception as error:
            raise ProcessingFailure(
                FailureReason.CORRUPT_DOCUMENT, detail=type(error).__name__
            ) from None
        return ExtractionResult(
            sections=builder.finish(), extractor=self.name, page_count=None, properties=properties
        )


def _check_archive(data: bytes, limits: ExtractionLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile:
        raise ProcessingFailure(FailureReason.CORRUPT_DOCUMENT, detail="not a zip") from None
    uncompressed = sum(entry.file_size for entry in entries)
    compressed = sum(entry.compress_size for entry in entries) or 1
    if (
        len(entries) > _MAX_ZIP_ENTRIES
        or uncompressed > limits.max_docx_uncompressed_bytes
        or uncompressed / compressed > _MAX_COMPRESSION_RATIO
    ):
        raise ProcessingFailure(FailureReason.CONTENT_TOO_LARGE, detail="archive limits")


def _heading_levels(document: Any) -> dict[str, int]:
    """Map style ids to heading levels once, instead of resolving a style per paragraph."""
    levels: dict[str, int] = {}
    for style in document.styles:
        name = style.name or ""
        if name == "Title":
            levels[style.style_id] = 1
        else:
            match = _HEADING_STYLE.match(name)
            if match:
                levels[style.style_id] = int(match.group(1))
    return levels


class _SectionBuilder:
    def __init__(self, limits: ExtractionLimits, heading_levels: dict[str, int]) -> None:
        self._limits = limits
        self._heading_levels = heading_levels
        self._total = 0
        self._sections: list[ExtractedSection] = []
        self._heading: str | None = None
        self._level: int | None = None
        self._lines: list[str] = []

    def _charge(self, text: str) -> None:
        self._total += len(text)
        if self._total > self._limits.max_text_chars:
            raise ProcessingFailure(FailureReason.CONTENT_TOO_LARGE)

    def _flush(self) -> None:
        if self._heading is None and not any(self._lines):
            return
        self._sections.append(
            ExtractedSection(
                kind=SectionKind.SECTION if self._heading is not None else SectionKind.BODY,
                text="\n".join(self._lines),
                heading=self._heading,
                heading_level=self._level,
            )
        )
        self._lines = []

    def add_paragraph(self, paragraph: Paragraph) -> None:
        text = paragraph.text
        self._charge(text)
        level = self._heading_levels.get(paragraph._p.style or "")
        if level is not None and text.strip():
            self._flush()
            self._heading, self._level = text.strip(), level
        else:
            self._lines.append(text)

    def add_table(self, table: Table) -> None:
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            line = " | ".join(cells)
            self._charge(line)
            self._lines.append(line)

    def finish(self) -> list[ExtractedSection]:
        self._flush()
        return self._sections


def _properties(document: Any) -> dict[str, Any]:
    core = document.core_properties
    found: dict[str, Any] = {}
    for attribute, key in (
        ("title", "title"),
        ("author", "author"),
        ("subject", "subject"),
        ("keywords", "keywords"),
    ):
        value = getattr(core, attribute, None)
        if isinstance(value, str) and value.strip():
            found[key] = value
    for attribute, key in (("created", "created"), ("modified", "modified")):
        value = getattr(core, attribute, None)
        if value is not None:
            found[key] = value.isoformat()
    return found
