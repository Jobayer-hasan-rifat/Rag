"""Baseline timings for the extraction pipeline (no database, broker or storage).

Run from `backend/` with the dev virtualenv:  python -m benchmarks.processing_baseline

Measures extraction and normalisation (the CPU-bound part of document processing) on
generated documents. Numbers depend on the machine; record them with the hardware.
"""

import io
import platform
import statistics
import sys
import time
import tracemalloc
from collections.abc import Callable
from dataclasses import dataclass

import docx

from app.core.documents.processing import Deadline, ExtractionLimits
from app.core.documents.scripts import profile_scripts
from app.parsers.registry import get_extractor
from app.services.processing_service import ProcessingService
from tests.pdf_factory import make_pdf

REPEATS = 5
LIMITS = ExtractionLimits(
    max_pages=2000, max_text_chars=20_000_000, max_docx_uncompressed_bytes=200 * 1024 * 1024
)
BANGLA = "বাংলাদেশের রাজধানী ঢাকা এবং এটি একটি পরীক্ষামূলক বাক্য যা বারবার ব্যবহৃত হয়। "
ENGLISH = "The quick brown fox jumps over the lazy dog while the benchmark runs repeatedly. "


@dataclass
class Case:
    name: str
    file_type: str
    data: bytes


def _pdf(pages: int, sentence: str) -> bytes:
    page = "\n".join((sentence * 2).strip() for _ in range(25))
    return make_pdf([page] * pages)


def _docx(paragraphs: int, sentence: str) -> bytes:
    document = docx.Document()
    for index in range(paragraphs):
        if index % 40 == 0:
            document.add_heading(f"Section {index // 40}", level=1)
        document.add_paragraph(sentence * 3)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _markdown(headings: int) -> bytes:
    return "".join(
        f"## Heading {i} {BANGLA[:12]}\n\n{ENGLISH * 3}\n\n- item {i}\n- item {i + 1}\n\n"
        for i in range(headings)
    ).encode()


def build_cases() -> list[Case]:
    return [
        Case("PDF 10 pages, English", "pdf", _pdf(10, ENGLISH)),
        Case("PDF 100 pages, English", "pdf", _pdf(100, ENGLISH)),
        Case("PDF 100 pages, Bangla", "pdf", _pdf(100, BANGLA)),
        Case("PDF 100 pages, mixed", "pdf", _pdf(100, BANGLA + ENGLISH)),
        Case("DOCX 2,000 paragraphs, mixed", "docx", _docx(2000, BANGLA + ENGLISH)),
        Case("TXT 5 MB, mixed", "txt", ((BANGLA + ENGLISH) * 40 + "\n").encode() * 450),
        Case("Markdown 2,000 headings", "md", _markdown(2000)),
    ]


def _time(call: Callable[[], object]) -> float:
    started = time.perf_counter()
    call()
    return time.perf_counter() - started


def measure(case: Case) -> dict[str, float | int | str]:
    extractor = get_extractor(case.file_type)
    extraction = extractor.extract(case.data, limits=LIMITS, deadline=Deadline(600))
    chars = sum(len(s.text) for s in extraction.sections)

    extract_times = [
        _time(lambda: extractor.extract(case.data, limits=LIMITS, deadline=Deadline(600)))
        for _ in range(REPEATS)
    ]
    deadline = Deadline(600)
    normalise_times = [
        _time(
            lambda: ProcessingService._normalise(
                extraction, extractor.collapse_inline_spaces, deadline
            )
        )
        for _ in range(REPEATS)
    ]
    sections = ProcessingService._normalise(extraction, extractor.collapse_inline_spaces, deadline)
    profile_time = _time(lambda: profile_scripts(s.text for s in sections))

    tracemalloc.start()
    extractor.extract(case.data, limits=LIMITS, deadline=Deadline(600))
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    extract_ms = statistics.median(extract_times) * 1000
    normalise_ms = statistics.median(normalise_times) * 1000
    total_ms = extract_ms + normalise_ms + profile_time * 1000
    return {
        "name": case.name,
        "file_kb": round(len(case.data) / 1024),
        "sections": len(extraction.sections),
        "chars": chars,
        "extract_ms": round(extract_ms, 1),
        "normalise_ms": round(normalise_ms, 1),
        "profile_ms": round(profile_time * 1000, 1),
        "total_ms": round(total_ms, 1),
        "mchars_per_s": round(chars / (total_ms / 1000) / 1e6, 2) if total_ms else 0,
        "py_peak_mb": round(peak / 1024 / 1024, 1),
    }


def main() -> None:
    print(f"Python {platform.python_version()} on {platform.platform()} ({platform.processor()})")
    import pymupdf

    print(f"PyMuPDF {pymupdf.__version__}; median of {REPEATS} runs")
    print("tracemalloc sees Python allocations only\n")
    columns = [
        "Document", "File KB", "Sections", "Chars", "Extract ms", "Normalise ms",
        "Profile ms", "Total ms", "M chars/s", "Py peak MB",
    ]  # fmt: skip
    print("| " + " | ".join(columns) + " |")
    print("|" + "---|" * len(columns))
    for case in build_cases():
        r = measure(case)
        cells = [
            r["name"], f"{r['file_kb']:,}", f"{r['sections']:,}", f"{r['chars']:,}",
            r["extract_ms"], r["normalise_ms"], r["profile_ms"], r["total_ms"],
            r["mchars_per_s"], r["py_peak_mb"],
        ]  # fmt: skip
        print("| " + " | ".join(str(cell) for cell in cells) + " |")
    sys.stdout.flush()


if __name__ == "__main__":
    main()
